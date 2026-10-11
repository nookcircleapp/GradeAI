"""Admin-only view of every paper's records in one place, and its Excel download."""
from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from app.database import SessionDep
from app.pilot.logic import final_score, is_flagged, max_score
from app.pilot.models import Paper, PaperSubmission, User
from app.pilot.schemas import AllRecords, RecordAnswer, RecordPaper, RecordRow
from app.pilot.security import AdminUser

router = APIRouter(prefix="/api/pilot/admin/records", tags=["pilot-admin"])


def collect(session: Session) -> AllRecords:
    papers = session.exec(
        select(Paper).where(Paper.is_template == False).order_by(Paper.created_at.desc())  # noqa: E712
    ).all()
    owners = {u.id: u for u in session.exec(select(User)).all()}
    subs = session.exec(select(PaperSubmission).order_by(PaperSubmission.started_at)).all()
    by_paper: dict[int, list[PaperSubmission]] = {}
    for s in subs:
        by_paper.setdefault(s.paper_id, []).append(s)

    out_papers, rows = [], []
    for paper in papers:
        owner = owners.get(paper.owner_id)
        out_papers.append(
            RecordPaper(
                id=paper.id,
                title=paper.title,
                subject=paper.subject,
                teacher=owner.name if owner else "",
                teacher_email=owner.email if owner else "",
                questions=[q["text"] for q in paper.questions],
                marks=[int(q["marks"]) for q in paper.questions],
                max_score=max_score(paper),
                created_at=paper.created_at,
            )
        )
        for s in by_paper.get(paper.id, []):
            answers = {a["question_index"]: a["answer"] for a in s.answers}
            grades = {g["question_index"]: g for g in (s.grades or [])}
            cells = []
            for i in range(len(paper.questions)):
                g = grades.get(i)
                override = s.overrides.get(str(i))
                cells.append(
                    RecordAnswer(
                        answer=answers.get(i, ""),
                        score=(override if override is not None else g["score"]) if g else None,
                        explanation=(g or {}).get("explanation") or "",
                    )
                )
            rows.append(
                RecordRow(
                    paper_id=paper.id,
                    submission_id=s.id,
                    student_name=s.student_name,
                    roll_number=s.roll_number,
                    section=s.section,
                    status=s.status,
                    submitted_at=s.submitted_at,
                    ai_score=s.ai_score,
                    final_score=final_score(s),
                    max_score=s.max_score,
                    flagged=is_flagged(s),
                    ai_fooled=s.ai_fooled,
                    answers=cells,
                )
            )
    return AllRecords(papers=out_papers, rows=rows)


@router.get("", response_model=AllRecords)
def all_records(_: AdminUser, session: SessionDep) -> AllRecords:
    return collect(session)


STATUS = {"in_progress": "Still writing", "grading": "Grading", "graded": "Graded", "failed": "Grading failed"}


def _cell(v):
    # Student text opened in Excel must not run as a formula.
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    return v


def _sheet_title(title: str, used: set[str]) -> str:
    base = "".join(c for c in title if c not in '[]:*?/\\').strip()[:28] or "Paper"
    name, n = base, 2
    while name.lower() in used:
        name = f"{base[:25]} ({n})"
        n += 1
    used.add(name.lower())
    return name


def _write(ws, header: list[str], rows: list[list], widths: dict[int, int]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([_cell(v) for v in row])
    ws.freeze_panes = "A2"
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 16)
    if rows:
        ws.auto_filter.ref = ws.dimensions


def workbook(data: AllRecords) -> bytes:
    wb = Workbook()
    papers = {p.id: p for p in data.papers}

    summary = wb.active
    summary.title = "All students"
    _write(
        summary,
        ["Paper", "Teacher", "Student", "Roll number", "Section", "Status", "Submitted", "AI score", "Final score", "Out of", "Percent", "AI flagged a trick"],
        [
            [
                papers[r.paper_id].title, papers[r.paper_id].teacher, r.student_name, r.roll_number, r.section,
                STATUS.get(r.status, r.status), r.submitted_at, r.ai_score, r.final_score, r.max_score,
                round(100 * r.final_score / r.max_score, 1) if r.final_score is not None and r.max_score else None,
                "Yes" if r.flagged else "",
            ]
            for r in data.rows
        ],
        {1: 30, 3: 24, 7: 20},
    )

    used = {"all students"}
    for paper in data.papers:
        ws = wb.create_sheet(_sheet_title(paper.title, used))
        header = ["Student", "Roll number", "Section", "Status", "Submitted", "Final score", "Out of"]
        for n, (text, marks) in enumerate(zip(paper.questions, paper.marks), start=1):
            short = text if len(text) <= 80 else text[:77] + "..."
            header += [f"Q{n} answer: {short}", f"Q{n} score (of {marks})", f"Q{n} AI reasoning"]
        rows = []
        for r in data.rows:
            if r.paper_id != paper.id:
                continue
            row = [r.student_name, r.roll_number, r.section, STATUS.get(r.status, r.status), r.submitted_at, r.final_score, r.max_score]
            for a in r.answers:
                row += [a.answer, a.score, a.explanation]
            rows.append(row)
        widths = {1: 24, 5: 20}
        for n in range(len(paper.questions)):
            widths[8 + 3 * n] = 50
            widths[10 + 3 * n] = 40
        _write(ws, header, rows, widths)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/export.xlsx")
def export_xlsx(_: AdminUser, session: SessionDep) -> Response:
    """Every paper's records: a summary sheet, then one sheet per paper with each answer."""
    return Response(
        content=workbook(collect(session)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="blinkscore-all-records.xlsx"'},
    )
